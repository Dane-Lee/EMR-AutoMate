"""Write a batch of physical assessments / follow-ups to an Excel workbook.

    python pa_workbook.py --demo        # fake data, proves the layout

ONE SHEET PER ASSESSMENT, laid out in the order the EMR asks for it (Dane, 2026-09-18):
page 2 fields, page 3 picks, page 4 in its four-section order, page 5 as the 1-9 answer
block with the backing text under it, page 6 last. He opens it on his own time and copies
each cell into the form, instead of being walked through six pages in chat.

WHAT THIS FILE DOES NOT DO
--------------------------
It is a formatter. Every rule about what may go in a cell lives in `pa_templates.md` and
is applied when the text is WRITTEN, not here:

  * Palpation and Observation carry Dane's own findings or a `[ bracket slot ]` -- never
    a plausible-sounding invention.
  * Red-flag sections list the items a scenario TOUCHES, for him to check. Nothing here
    writes a red-flag field as negative and nothing clears anyone.
  * Protective recommendations are permissive sentences, never restrictions.

NO PHI. The sheets carry scenarios and field text, never a name, date of birth or
identifier -- he matches a sheet to a person by its ref, on his side. The output is
gitignored all the same, because a day's real assessments is his working file.
"""

import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HERE = os.path.dirname(os.path.abspath(__file__))

HEAD_FILL = PatternFill("solid", fgColor="1F3864")      # section bar
HEAD_FONT = Font(bold=True, color="FFFFFF", size=12)
FIELD_FONT = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

# The page-5 checklist, in the on-screen order Dane corrected on 2026-09-11: the
# protective-recommendations conditional sits after item 7, before First Aid.
COACHING_DETAILS = [
    "Ergonomic work station adjustments",
    "Job coaching specific to behaviors, postures noted above",
    "ADL coaching specific to risks listed above",
    "Reminded the employee of preventative job-specific mobility program",
    "Reminded the employee of preventative job-specific conditioning activities",
    "Educated employee on total health and wellness",
    "The employee is capable of performing all essential job tasks",
    "Requires protective recommendations",
    "First Aid provided",
]

# Page 4, in the form's order -- not the reference sheet's (Dane, 2026-09-11).
RED_FLAG_SECTIONS = [("signs", "1. Red Flag Signs"),
                     ("mechanism", "2. Red Flag Mechanism of Injury"),
                     ("symptoms", "3. Red Flag Symptoms"),
                     ("personal", "4. Red Flag Personal Conditions")]

PAGE3_FIELDS = [("work_factors", "Work factors"),
                ("adls", "ADL's"),
                ("postures", "Postures"),
                ("faulty", "Faulty behaviors"),
                ("movement", "Suboptimal Human Movement Patterns")]


def _sheet_title(rec, used):
    """A tab name Excel will accept: 31 chars, none of []:*?/\\, and unique."""
    raw = f"{rec['ref']} {rec.get('title', '')}".strip()
    for ch in "[]:*?/\\":
        raw = raw.replace(ch, " ")
    raw = " ".join(raw.split())[:31] or f"Ref {rec['ref']}"
    title, n = raw, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = raw[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


class _Sheet:
    """Two columns: the field name, and the text to copy out of."""

    def __init__(self, ws):
        self.ws = ws
        self.row = 1
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 104

    def section(self, text):
        self.blank()
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        bar = self.ws.cell(row=self.row, column=2)
        bar.fill = HEAD_FILL
        self.row += 1

    def field(self, label, value, wrap=True):
        if value is None or value == "" or value == []:
            return
        if isinstance(value, (list, tuple)):
            value = "  ·  ".join(str(v) for v in value)
        self.ws.cell(row=self.row, column=1, value=label).font = FIELD_FONT
        self.ws.cell(row=self.row, column=1).alignment = TOP
        cell = self.ws.cell(row=self.row, column=2, value=str(value))
        cell.alignment = WRAP if wrap else TOP
        self.row += 1

    def note(self, text):
        cell = self.ws.cell(row=self.row, column=2, value=text)
        cell.alignment = WRAP
        cell.font = Font(italic=True, color="666666", size=10)
        self.row += 1

    def blank(self):
        self.row += 1


def _write_assessment(ws, rec):
    sh = _Sheet(ws)
    follow_up = rec.get("kind", "initial").lower().startswith("follow")

    kind = "Follow-Up Assessment" if follow_up else "Initial Physical Assessment"
    head = f"#{rec['ref']}   ·   {kind}"
    if rec.get("bucket"):
        head += f"   ·   {rec['bucket']}"
    cell = ws.cell(row=sh.row, column=1, value=head)
    cell.font = Font(bold=True, size=14)
    sh.row += 1
    if rec.get("title"):
        sh.note(rec["title"])

    # ── page 2 ──
    p2 = rec.get("page2", {})
    sh.section("PAGE 2 — Symptom Details")
    sh.field("Incident Details", p2.get("incident"))
    if p2.get("incident_long"):
        sh.field("Incident Details (longer)", p2["incident_long"])
    if not follow_up:
        # Dropped on a follow-up: both were settled on the initial assessment.
        sh.field("Primary Complaint", p2.get("primary_complaint"), wrap=False)
        sh.field("Categorize Mechanism", p2.get("mechanism"), wrap=False)
    sh.field("Palpation Comments", p2.get("palpation"))
    sh.field("Observation Comments", p2.get("observation"))

    # ── page 3 ──
    p3 = rec.get("page3", {})
    if p3:
        sh.section("PAGE 3 — Root Cause Analysis")
        for key, label in PAGE3_FIELDS:
            sh.field(label, p3.get(key), wrap=False)
        sh.field("Functionality", p3.get("functionality"))
        sh.field("Symptom Self-Management", p3.get("self_management"), wrap=False)
        sh.field("… Other (free text)", p3.get("self_management_other"))

    # ── page 4 ──
    p4 = rec.get("page4", {})
    if p4:
        sh.section("PAGE 4 — Red Flag Triage")
        sh.note("Matches surfaced against the reference sheet, for you to check or not. "
                "Two or more is the referral threshold.")
        for key, label in RED_FLAG_SECTIONS:
            items = p4.get(key)
            sh.field(label, items if items else "nothing surfaced", wrap=bool(items))
        if p4.get("count") is not None:
            sh.field("Count", p4["count"], wrap=False)
        for key, label in RED_FLAG_SECTIONS:
            sh.field(f"{label.split('. ', 1)[1]} — explanation",
                     p4.get(f"{key}_text"))

    # ── page 5 ──
    p5 = rec.get("page5", {})
    if p5:
        sh.section("PAGE 5 — Corrective Actions")
        answers = p5.get("answers", {})
        for i, item in enumerate(COACHING_DETAILS, start=1):
            sh.field(f"{i}. {item}", answers.get(item), wrap=False)
        contribution = p5.get("contribution", {})
        if contribution and not follow_up:
            sh.field("Contribution Type", contribution.get("bucket"), wrap=False)
            for label, value in contribution.get("risks", []):
                sh.field(f"… {label}", value, wrap=False)
        sh.blank()
        for label, text in p5.get("texts", []):
            sh.field(label, text)

    # ── page 6 ──
    p6 = rec.get("page6", {})
    if p6:
        sh.section("PAGE 6 — Plan")
        sh.field("Initial Plan", p6.get("initial_plan"), wrap=False)
        sh.field("Assessment Recommendation", p6.get("recommendation"), wrap=False)
        sh.field("Additional Comments", p6.get("comments"))
        for n, option in enumerate(p6.get("focus", []), start=1):
            sh.field(f"Focus for Next Encounter — option {n}", option)
        sh.field("Assessment Type", p6.get("assessment_type"), wrap=False)
        sh.field("Date of Next Encounter", p6.get("next_date"), wrap=False)

    ws.freeze_panes = "A2"


def _write_coaching(ws, rec):
    """A coaching encounter: one description, its type, and the boxes to tick.

    Same workbook as the assessments (Dane, 2026-09-21) because a morning's drafts are a
    mix of both, and hunting through two files to finish one day's work is the thing this
    was built to stop.
    """
    sh = _Sheet(ws)
    cell = ws.cell(row=sh.row, column=1,
                   value=f"#{rec['ref']}   ·   Coaching Encounter")
    cell.font = Font(bold=True, size=14)
    sh.row += 1
    if rec.get("title"):
        sh.note(rec["title"])

    sh.section("COACHING ENCOUNTER")
    sh.field("Coaching type", rec.get("coaching_type"), wrap=False)
    sh.field("Details", rec.get("details"), wrap=False)
    sh.field("Description", rec.get("description"))
    if rec.get("description_alt"):
        sh.field("Description (alternate)", rec["description_alt"])
    sh.field("Notes", rec.get("notes"))
    ws.freeze_panes = "A2"


def _write_index(ws, records):
    ws.title = "Index"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 70
    for col, label in enumerate(("Ref", "Type", "What it is"), start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for i, rec in enumerate(records, start=2):
        if rec.get("kind", "").lower() == "coaching":
            kind = "Coaching"
        elif rec.get("kind", "").lower().startswith("follow"):
            kind = "Follow-Up"
        else:
            kind = "Initial"
        if rec.get("bucket"):
            kind += f" · {rec['bucket']}"
        ws.cell(row=i, column=1, value=str(rec["ref"])).alignment = TOP
        ws.cell(row=i, column=2, value=kind).alignment = TOP
        ws.cell(row=i, column=3, value=rec.get("title", "")).alignment = WRAP
    ws.freeze_panes = "A2"


def write_workbook(records, path=None):
    """Write one sheet per assessment plus an index. Returns the path."""
    if not records:
        raise ValueError("nothing to write")
    path = path or os.path.join(
        _HERE, f"pa_worksheets_{date.today().strftime('%Y-%m-%d')}.xlsx")
    wb = Workbook()
    _write_index(wb.active, records)
    used = set()
    for rec in records:
        ws = wb.create_sheet(_sheet_title(rec, used))
        if rec.get("kind", "").lower() == "coaching":
            _write_coaching(ws, rec)
        else:
            _write_assessment(ws, rec)
    wb.save(path)
    return path


def _demo_records():
    """Fake data — the shape, not anyone's assessment."""
    return [
        {"ref": "1", "kind": "initial", "bucket": "NWC",
         "title": "Rolled ankle, off-the-job recreation",
         "page2": {"incident": "EE said they rolled their ankle over the weekend "
                               "playing a recreational sport. EE said it did not "
                               "happen at work.",
                   "primary_complaint": "musculoskeletal acute",
                   "mechanism": "slip/trip/fall  ·  or other",
                   "palpation": "Tenderness present over [ lateral ankle ].",
                   "observation": "EE experienced discomfort with single-leg stance."},
         "page3": {"work_factors": ["No known work factors at this time"],
                   "adls": ["Sports and/or hobbies"],
                   "postures": ["No observed postures at this time"],
                   "faulty": ["No faulty behaviors observed at this time"],
                   "movement": ["Stiff or limited ankle mobility"],
                   "functionality": "Ankle.",
                   "self_management": ["Ice", "Support"]},
         "page4": {"signs": ["Guarded / loss of ROM"],
                   "mechanism": [], "symptoms": [], "personal": [], "count": 1,
                   "signs_text": "Limited ankle ROM with tenderness laterally."},
         "page5": {"answers": {"Ergonomic work station adjustments": "No",
                               "Job coaching specific to behaviors, postures noted above": "No",
                               "ADL coaching specific to risks listed above": "Yes",
                               "Reminded the employee of preventative job-specific mobility program": "Yes",
                               "Reminded the employee of preventative job-specific conditioning activities": "No",
                               "Educated employee on total health and wellness": "Yes",
                               "The employee is capable of performing all essential job tasks": "Yes",
                               "Requires protective recommendations": "Yes",
                               "First Aid provided": "No"},
                   "contribution": {"bucket": "NWC",
                                    "risks": [("Risk of progressing to an injury or illness", "Low"),
                                              ("Risk of becoming work-related", "Low"),
                                              ("ATI licensed provider providing MTBFA", "No")]},
                   "texts": [("ADL coaching specific to risks listed above",
                              "EIS and EE discussed easing off recreational activity "
                              "while the ankle settles down."),
                             ("Protective recommendations",
                              "Encourage attention to footing and periodic position "
                              "changes as workflow allows.")]},
         "page6": {"initial_plan": "Continue Work",
                   "recommendation": "none",
                   "focus": ["Recheck ankle ROM and tenderness.",
                             "Check whether the ice and support helped."],
                   "assessment_type": "PA Follow up"}},
        {"ref": "2", "kind": "follow-up", "bucket": "NWC",
         "title": "Upper back, improving",
         "page2": {"incident": "EE said the upper back has improved but is still "
                               "present.",
                   "observation": "Reduced tension through the upper back."},
         "page5": {"answers": {"Educated employee on total health and wellness": "Yes",
                               "The employee is capable of performing all essential job tasks": "Yes",
                               "Requires protective recommendations": "No",
                               "First Aid provided": "No"},
                   "texts": [("Educated employee on total health and wellness",
                              "EIS and EE discussed what could be done over the next "
                              "few days to keep the upper back improving.")]},
         "page6": {"initial_plan": "Continue Work",
                   "focus": ["Recheck whether the tightness has kept easing."],
                   "assessment_type": "PA Follow up"}},
    ]


if __name__ == "__main__":
    if "--demo" in sys.argv:
        out = write_workbook(_demo_records(),
                             os.path.join(_HERE, "debug", "pa_worksheets_DEMO.xlsx"))
        print(f"wrote {out}")
    else:
        print(__doc__.strip().splitlines()[0])
        print("\nThis module is written to; import write_workbook(records, path).")
