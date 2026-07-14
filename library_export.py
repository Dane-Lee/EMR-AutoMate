"""
library_export.py — make the description library readable by Copilot
====================================================================
Copilot (in VS Code) can read text files, but not a binary .xlsx. This dumps
`EMR Easy Enter Worksheets.xlsx` — the library of standard, reusable encounter
descriptions — into `description_library.md`, so Copilot can pull an existing
description instead of inventing a new one every time.

RUN IT
    python library_export.py

    Re-run whenever you add descriptions to the worksheet. Copilot can also run it
    itself if the .md looks stale.

OUTPUT IS TREATED AS SENSITIVE
    description_library.md is gitignored, exactly like the .xlsx it comes from. The
    worksheet is meant to hold blank templates with ____ placeholders, but it lives
    in the PHI side of the fence and is not worth the argument — Copilot is cleared
    for PHI and is the only assistant that reads it.

    This script deliberately prints only counts, never cell contents, so its own
    output is safe to show anyone.
"""

import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed.  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_XLSX = os.path.join(HERE, "EMR Easy Enter Worksheets.xlsx")
OUTPUT_MD = os.path.join(HERE, "description_library.md")


def cell(v):
    """Render a cell for a markdown table: never None, never a broken pipe."""
    if v is None:
        return ""
    return str(v).replace("|", "\\|").replace("\n", " ").strip()


def export():
    if not os.path.exists(SOURCE_XLSX):
        sys.exit(f"Not found: {os.path.basename(SOURCE_XLSX)}\n"
                 "Put the worksheet in this folder and run again.")

    wb = load_workbook(SOURCE_XLSX, data_only=True, read_only=True)

    out = [
        "# Standard Description Library",
        "",
        "Auto-generated from `EMR Easy Enter Worksheets.xlsx` by `library_export.py`.",
        "**Do not edit by hand** — edit the workbook and re-run the script.",
        "",
        "Reuse a description from here whenever one fits the dictated encounter,",
        "filling any `____` placeholder with the specifics that were actually said.",
        "Keep the third-person EIS/EE voice. Only write a brand-new description when",
        "nothing here fits.",
        "",
    ]

    sheet_count = 0
    row_total = 0

    for name in wb.sheetnames:
        ws = wb[name]
        rows = [r for r in ws.iter_rows(values_only=True)
                if r and any(c is not None and str(c).strip() for c in r)]
        if not rows:
            continue

        sheet_count += 1
        # First non-empty row is the header; the rest are descriptions.
        header, body = rows[0], rows[1:]
        width = max(len(r) for r in rows)

        out.append(f"## {name}")
        out.append("")
        head = [cell(header[i]) if i < len(header) else "" for i in range(width)]
        out.append("| " + " | ".join(h or f"col{i+1}" for i, h in enumerate(head)) + " |")
        out.append("|" + "|".join(["---"] * width) + "|")
        for r in body:
            cells = [cell(r[i]) if i < len(r) else "" for i in range(width)]
            out.append("| " + " | ".join(cells) + " |")
            row_total += 1
        out.append("")

    with open(OUTPUT_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))

    # Counts only — no cell contents — so this output is safe to paste anywhere.
    print(f"Wrote {os.path.basename(OUTPUT_MD)}")
    print(f"  sheets exported : {sheet_count}")
    print(f"  description rows: {row_total}")
    print("\nOpen description_library.md to review it. Copilot reads it from now on.")


if __name__ == "__main__":
    export()
