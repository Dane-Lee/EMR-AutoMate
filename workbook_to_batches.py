"""Turn Dane's own encounter-tracking workbook into builder-ready dictation batches.

    python workbook_to_batches.py                 # write the batch files
    python workbook_to_batches.py --check         # validate only, write nothing

Dane tracks a week of encounters in his own spreadsheet (`Encounter Work/`), one tab per
day, and splits a day into `9-23` / `9-23 (B)` when the same person needs two encounters
— the EMR disables "+ Add Case" for anyone who already has an In Progress case, so the
second one has to wait for the first to be finalized. **One tab is one run.**

His column vocabulary is his own: `H&W`, `Group`, `Tool Adjustment`, `Ergonimic`. The EMR
takes none of those. This module is the translation layer, and it is deliberately a
LOOKUP, not a guess:

  * `TYPE` maps his Type column onto `ace.COACHING_TYPE_UUIDS`. An unmapped value is
    REPORTED AND DROPPED, never approximated — an off-list coaching type is refused by
    the builder anyway, and a silently "corrected" one is a controlled value in a record
    that Dane did not choose.
  * `DETAIL` maps his Optional column onto `ace.CHECKBOX_UUID_MAP` for that type. The
    spellings really do differ: he writes `Tool Adjustment`, the EMR wants
    `Tools adjustment`, and `load_dictated_batch` drops anything that does not match.
  * Assessment rows (`F/U`, `F/U Assessment`, `Initial PA`) are carried straight through
    to `assessment_rows()` and never enter a batch. The entry engine only ever clicks the
    Coaching Encounter tile; there is no automated path for an assessment.

The descriptions are NOT generated here. They are written once, reviewed, and stored in
`Encounter Work/descriptions.json` keyed `tab|code`; this module only joins them on.

NO NAMES. The `Code` column is the ref, and the code -> person mapping lives in Dane's
cheat sheet, which the builder reads locally (`encounter_builder.load_code_map`). Nothing
in the output carries a name, and `Location` is deliberately dropped: department,
division and shift come off the roster, which is the only place they are trustworthy.
"""

import argparse
import glob
import json
import os
import sys

import ati_coaching_encounter as ace

_HERE = os.path.dirname(os.path.abspath(__file__))
WORK_DIR = os.path.join(_HERE, "Encounter Work")
DESCRIPTIONS = os.path.join(WORK_DIR, "descriptions.json")
BATCH_DIR = os.path.join(WORK_DIR, "batches")

# His Type column -> the EMR's coaching types. Lowercase keys; the two spellings of
# "Ergonomic" are both his and both mean the same thing.
TYPE = {
    "h&w": "Health/Wellness Coaching",
    "group": "Group Class",
    "ergonomic adjustment": "Ergonomic Adjustment",
    "ergonimic adjustment": "Ergonomic Adjustment",
    "safety": "Safety Coaching",
    "relationship development": "Relationship Development Encounter",
    "near miss education": "Near Miss Education",
    "job specific coaching": "Job-Specific Coaching",
}

# Assessment types. Not a coaching type, no automated path, handled by hand.
ASSESSMENT = {"f/u", "f/u assessment", "initial pa", "pa", "hma", "hma f/u",
              "task assessment", "physical assessment", "follow-up assessment"}

# His Optional column -> real checkbox labels. Keys are lowercase fragments as he writes
# them; the values are verbatim from ace.CHECKBOX_UUID_MAP and must stay that way.
DETAIL = {
    "slip/trip/fall prevention": "Slip/trip/fall Prevention",
    "slip/trip/fall prevention ": "Slip/trip/fall Prevention",
    "unsafe behavior": "Unsafe Behavior",
    "safety hazard": "Safety Hazard",
    "ppe use": "PPE Use",
    "3-point contact": "3-point contact",
    "awareness/alertness": "Awareness/alertness",
    "tool adjustment": "Tools adjustment",
    "tools adjustment": "Tools adjustment",
    "industrial adjustment": "Industrial ergo adjustment",
    "office adjustment": "Office ergo adjustment",
    "group prevention mobility": "Group Preventative Mobility",
    "group preventative mobility": "Group Preventative Mobility",
    "pre-shift meeting": "Pre-shift Meeting",
    "management meetings": "Management Meetings",
    "new hire orientation": "New Hire Orientation",
    "safety meeting": "Safety Meeting",
    "hydration": "Hydration",
    "nutrition": "Nutrition",
    "execise": "Personal exercise/fitness",          # his spelling
    "exercise": "Personal exercise/fitness",
    "personal exercise/fitness": "Personal exercise/fitness",
    "sleep": "Sleep",
    "stress": "Stress",
    "current employee": "Current Employee",
    "new hire": "New Hire",
    "body mechanics": "Body Mechanics",
    "material handling": "Material handling",
    "proper lifting": "Proper lifting",
    "proper loading/unloading": "Proper loading/unloading",
    "proper push/pull": "Proper push/pull",
    "postural/position coaching": "Postural/Position Coaching",
    "tool/equipment handling": "Tool/equipment handling",
    "other job coaching": "Other job coaching",
    "rest (task design)": "Rest (task design)",
    "rest break": "Rest break",
    "other": "Other",
    "none": None,                                    # he writes this for "no boxes"
}


def _workbook_path():
    """The newest tracking workbook in Encounter Work/. Explicit about finding none."""
    hits = sorted(glob.glob(os.path.join(WORK_DIR, "*.xlsx")),
                  key=os.path.getmtime, reverse=True)
    hits = [h for h in hits if not os.path.basename(h).startswith("~$")]
    if not hits:
        raise SystemExit(f"No .xlsx found in {WORK_DIR}")
    return hits[0]


def _split_details(raw, ctype):
    """His Optional cell -> (valid checkbox labels, unrecognised fragments).

    Split on commas, look each fragment up, then keep only what this coaching type
    actually offers. A fragment that maps to a real label but not for THIS type is an
    unrecognised one, not a silent drop: Near Miss Education has no boxes at all.
    """
    valid = ace.CHECKBOX_UUID_MAP.get(ctype, {})
    keep, unknown = [], []
    for frag in str(raw or "").split(","):
        frag = frag.strip()
        if not frag:
            continue
        label = DETAIL.get(frag.lower(), "__MISS__")
        if label is None:                      # "none" — he means no boxes
            continue
        if label == "__MISS__":
            unknown.append(frag)
            continue
        match = next((v for v in valid if v.lower() == label.lower()), None)
        if match is None:
            unknown.append(frag)
        elif match not in keep:
            keep.append(match)
    return keep, unknown


def read_rows(path=None):
    """Every row of the workbook, classified. Returns (coaching, assessment, problems).

    `coaching` is {tab: [row, ...]} in sheet order — the tab IS the batch, because that
    is how Dane split them to avoid two drafts for one person at once.
    """
    import openpyxl
    path = path or _workbook_path()
    wb = openpyxl.load_workbook(path, data_only=True)
    coaching, assessment, problems = {}, [], []

    for ws in wb.worksheets:
        hdr = {str(c.value).strip().lower(): c.column - 1 for c in ws[1] if c.value}
        if not all(k in hdr for k in ("#", "code", "type", "description")):
            problems.append((ws.title, "headers don't match, tab skipped"))
            continue
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            def cell(key):
                i = hdr.get(key)
                if i is None or i >= len(r) or r[i] in (None, ""):
                    return ""
                return str(r[i]).strip()

            code, raw_type = cell("code"), cell("type")
            if not raw_type:
                continue
            low = raw_type.lower()
            rec = {"tab": ws.title, "n": cell("#"), "code": code,
                   "raw_type": raw_type, "optional": cell("optional"),
                   "location": cell("location"), "raw": cell("description")}
            if low in ASSESSMENT:
                assessment.append(rec)
                continue
            ctype = TYPE.get(low)
            if not ctype:
                problems.append((f"{ws.title}/{code}", f"unmapped type {raw_type!r}"))
                continue
            details, unknown = _split_details(rec["optional"], ctype)
            for u in unknown:
                problems.append((f"{ws.title}/{code}",
                                 f"detail {u!r} is not an option for {ctype}"))
            rec.update({"coaching_type": ctype, "details": details})
            rows.append(rec)
        if rows:
            coaching[ws.title] = rows
    return coaching, assessment, problems


def build_batches(coaching, descriptions):
    """{tab: batch dict} in dictated_batch.json shape. Missing description = dropped."""
    from datetime import date
    batches, missing = {}, []
    for tab, rows in coaching.items():
        out = []
        for rec in rows:
            key = f"{tab}|{rec['code']}"
            entry = descriptions.get(key)
            if not entry or not entry[0]:
                missing.append(key)
                continue
            text, library = (entry + [None])[:2] if isinstance(entry, list) else (entry, None)
            out.append({"ref": rec["code"],
                        "coaching_type": rec["coaching_type"],
                        "details": rec["details"],
                        "description": text,
                        "library": library or ""})
        if out:
            batches[tab] = {"written": date.today().isoformat(),
                            "source_tab": tab, "rows": out}
    return batches, missing


def assessment_rows(assessment):
    """The rows that must be hand-entered, in tab order. Nothing is invented here."""
    return sorted(assessment, key=lambda r: (r["tab"], int(r["n"] or 0)))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--check", action="store_true", help="validate only, write nothing")
    args = ap.parse_args(argv)

    path = _workbook_path()
    print(f"workbook: {os.path.basename(path)}")
    coaching, assessment, problems = read_rows(path)

    with open(DESCRIPTIONS, encoding="utf-8") as fh:
        descriptions = {k: v for k, v in json.load(fh).items() if not k.startswith("_")}

    batches, missing = build_batches(coaching, descriptions)

    n_coach = sum(len(v) for v in coaching.values())
    n_out = sum(len(b["rows"]) for b in batches.values())
    print(f"  coaching rows: {n_coach}  ->  {n_out} in {len(batches)} batch file(s)")
    print(f"  assessment rows (hand entry): {len(assessment)}")
    for where, why in problems:
        print(f"  ! {where}: {why}")
    for key in missing:
        print(f"  ! {key}: no description written yet")

    if args.check:
        print("--check: nothing written.")
        return 0 if not (problems or missing) else 1

    os.makedirs(BATCH_DIR, exist_ok=True)
    for i, (tab, batch) in enumerate(batches.items(), start=1):
        safe = tab.replace(" ", "_").replace("(", "").replace(")", "")
        dest = os.path.join(BATCH_DIR, f"{i:02d}_{safe}.json")
        with open(dest, "w", encoding="utf-8") as fh:
            json.dump(batch, fh, indent=1, ensure_ascii=False)
        print(f"  wrote {os.path.relpath(dest, _HERE)}  ({len(batch['rows'])} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
